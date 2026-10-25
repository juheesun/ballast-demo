"""과거 이력 API (MVP4)

- data/dataset.json (실제 10 + 시나리오 96) 은 읽기만 한다. 기존 Backend.history() 를 그대로 쓴다.
- 엑셀(.xlsx)로 올린 새 입찰 기록은 SQLite 파일에 영구 저장한다 (origin = "upload").
  기본 위치 data/history.db · 테스트할 때는 환경변수 BALLAST_HISTORY_DB 로 다른 파일을 지정한다.
- 통계(낙찰률·실주 사유 등)는 화면(HistoryPage.jsx)에서만 계산한다. 서버는 행 목록만 준다.
"""
from __future__ import annotations
import io, os, sqlite3
from datetime import date, datetime
from pathlib import Path

from fastapi import APIRouter, HTTPException, Query, Request

ROOT = Path(__file__).resolve().parents[2]
DB_PATH = Path(os.environ.get("BALLAST_HISTORY_DB", ROOT / "data" / "history.db"))
MAX_BYTES = 5 * 1024 * 1024
LOSS_REASONS = ["가격", "납기", "사양", "금융", "기타"]

# 엑셀 첫 줄(제목 줄)의 이름 → 내부 필드. 필수: 입찰일·조선소·견적·시장가·결과
COLUMNS = {"입찰일": "date", "조선소": "shipyard", "선주": "owner", "척수": "ships", "견적($M)": "quote",
           "시장가($M)": "market", "결과": "result", "실주 사유": "loss_reason", "마진(%)": "margin", "메모": "note"}
REQUIRED = ["입찰일", "조선소", "견적($M)", "시장가($M)", "결과"]

router = APIRouter(prefix="/api/history", tags=["history"])
be = None                                 # main.py 가 Backend 를 넣어 준다 (dataset.json 행을 읽는 데 씀)

SCHEMA = """
CREATE TABLE IF NOT EXISTS bid_history (
    id          INTEGER PRIMARY KEY AUTOINCREMENT,
    bid_date    TEXT    NOT NULL,              -- YYYY-MM-DD
    shipyard    TEXT    NOT NULL,
    owner       TEXT    NOT NULL DEFAULT '',
    num_ships   INTEGER,
    quote_usd_mn  REAL  NOT NULL CHECK (quote_usd_mn > 0),    -- 척당 견적 $M
    market_usd_mn REAL  NOT NULL CHECK (market_usd_mn > 0),   -- 척당 시장 기준가 $M
    won         INTEGER NOT NULL CHECK (won IN (0, 1)),
    loss_reason TEXT,                          -- 실주일 때만
    margin      REAL,                          -- 입력값(소수). 출처를 확인하지 않은 값이라 '추정'으로 표시
    note        TEXT    NOT NULL DEFAULT '',
    file_name   TEXT    NOT NULL DEFAULT '',   -- 올린 엑셀 파일 이름
    created_at  TEXT    NOT NULL
)"""


def _db() -> sqlite3.Connection:
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    con = sqlite3.connect(DB_PATH)
    con.row_factory = sqlite3.Row
    con.execute(SCHEMA)
    return con


def _row_out(r: sqlite3.Row) -> dict:
    """dataset.json 행과 같은 모양으로 맞춘다 → 화면의 필터·통계를 그대로 쓴다."""
    q, m = r["quote_usd_mn"], r["market_usd_mn"]
    return {"id": f"UP-{r['id']:04d}", "anchor": None, "origin": "upload", "variant": "UPLOAD",
            "date": r["bid_date"], "shipyard": r["shipyard"], "owner": r["owner"], "ships": r["num_ships"],
            "delivery": None, "quote": round(q, 1), "market": round(m, 1), "ratio": round(q / m, 4),
            "won": bool(r["won"]), "loss_reason": r["loss_reason"], "margin": r["margin"],
            "note": r["note"], "source": f"엑셀 업로드: {r['file_name']} ({r['created_at'][:10]})"}


def _key(date_, shipyard, owner, quote, market, won) -> tuple:
    """중복 판단: 입찰일·조선소·선주·견적·시장가·결과가 모두 같으면 같은 기록."""
    return (date_, shipyard, owner, round(float(quote), 2), round(float(market), 2), bool(won))


def _mark_duplicates(con: sqlite3.Connection, rows: list[dict]) -> None:
    """오류 없는 줄에 duplicate 사유를 붙인다. 중복 줄은 저장할 때 건너뛴다 (오류는 아님)."""
    saved = {_key(r["bid_date"], r["shipyard"], r["owner"], r["quote_usd_mn"], r["market_usd_mn"], r["won"])
             for r in con.execute("SELECT * FROM bid_history")}
    seen = set()
    for r in rows:
        r["duplicate"] = None
        if r["errors"]:
            continue
        k = _key(r["date"], r["shipyard"], r["owner"], r["quote"], r["market"], r["won"])
        if k in saved:
            r["duplicate"] = "이미 저장된 기록"
        elif k in seen:
            r["duplicate"] = "파일 안에서 겹치는 줄"
        seen.add(k)


def _parse(content: bytes) -> tuple[list[dict], list[str]]:
    """엑셀 → (행 목록, 파일 전체 오류). 각 행에는 errors 목록이 붙는다."""
    from openpyxl import load_workbook
    try:
        ws = load_workbook(io.BytesIO(content), read_only=True, data_only=True).worksheets[0]
    except Exception:
        return [], ["엑셀(.xlsx) 파일을 열 수 없습니다."]
    lines = ws.iter_rows(values_only=True)
    head = [str(h).strip() if h is not None else "" for h in next(lines, [])]
    missing = [c for c in REQUIRED if c not in head]
    if missing:
        return [], [f"첫 줄에 필요한 열이 없습니다: {', '.join(missing)}"]
    idx = {COLUMNS[h]: i for i, h in enumerate(head) if h in COLUMNS}

    rows = []
    for n, line in enumerate(lines, start=2):
        cell = lambda k: line[idx[k]] if k in idx and idx[k] < len(line) else None
        if all(v is None or str(v).strip() == "" for v in line):
            continue
        errs, r = [], {"line": n}

        d = cell("date")
        if isinstance(d, (datetime, date)):
            r["date"] = d.strftime("%Y-%m-%d")
        else:
            try:
                r["date"] = datetime.strptime(str(d).strip()[:10], "%Y-%m-%d").strftime("%Y-%m-%d")
            except (ValueError, TypeError):
                errs.append("입찰일 형식 오류 (예: 2025-04-28)")
                r["date"] = None if d is None else str(d)

        r["shipyard"] = str(cell("shipyard") or "").strip()
        if not r["shipyard"]:
            errs.append("조선소가 비어 있음")
        r["owner"] = str(cell("owner") or "").strip()
        r["note"] = str(cell("note") or "").strip()

        for k, label in (("quote", "견적"), ("market", "시장가")):
            try:
                r[k] = float(cell(k))
                if r[k] <= 0:
                    raise ValueError
            except (TypeError, ValueError):
                errs.append(f"{label}은 0보다 큰 숫자여야 함")
                r[k] = None

        s = cell("ships")
        try:
            r["ships"] = None if s in (None, "") else int(s)
        except (TypeError, ValueError):
            errs.append("척수는 정수여야 함")
            r["ships"] = None

        res = str(cell("result") or "").strip()
        r["won"] = {"낙찰": True, "실주": False}.get(res)
        if r["won"] is None:
            errs.append("결과는 '낙찰' 또는 '실주'")
        reason = str(cell("loss_reason") or "").strip() or None
        if r["won"] is False and reason not in LOSS_REASONS:
            errs.append(f"실주 사유는 {'/'.join(LOSS_REASONS)} 중 하나")
        r["loss_reason"] = reason if r["won"] is False else None

        mg = cell("margin")
        try:                                          # 8.5 또는 '8.5%' → 0.085
            r["margin"] = None if mg in (None, "") else round(float(str(mg).replace("%", "")) / 100, 4)
        except ValueError:
            errs.append("마진(%)은 숫자여야 함")
            r["margin"] = None

        r["errors"] = errs
        rows.append(r)
    if not rows:
        return [], ["데이터 줄이 없습니다."]
    return rows, []


@router.get("")
def history():
    """dataset.json 행 + 업로드 행. 응답 모양은 기존과 같고 rows 뒤에 업로드분이 붙는다."""
    base = be.history()
    with _db() as con:
        up = [_row_out(r) for r in con.execute("SELECT * FROM bid_history ORDER BY bid_date, id")]
    return {**base, "rows": base["rows"] + up, "uploaded": len(up)}


@router.post("/upload")
async def upload(request: Request, save: bool = Query(False, description="false = 미리보기만, true = 저장"),
                 name: str = Query("", max_length=200, description="파일 이름 (기록용)")):
    """요청 본문에 .xlsx 파일을 그대로 보낸다. 오류가 한 줄이라도 있으면 저장하지 않는다.
    이미 저장된 기록·파일 안에서 겹치는 줄은 저장하지 않고 건너뛴다 (skipped)."""
    content = await request.body()
    if not content:
        raise HTTPException(400, "파일이 비어 있습니다.")
    if len(content) > MAX_BYTES:
        raise HTTPException(413, "파일은 5MB 이하만 올릴 수 있습니다.")
    rows, file_errors = _parse(content)
    n_err = sum(bool(r["errors"]) for r in rows)
    with _db() as con:
        _mark_duplicates(con, rows)
        new = [r for r in rows if not r["errors"] and not r["duplicate"]]
        out = {"rows": rows, "file_errors": file_errors, "error_rows": n_err,
               "new_rows": len(new), "skipped": sum(bool(r["duplicate"]) for r in rows), "saved": 0}
        if save:
            if file_errors or n_err:
                raise HTTPException(422, "오류가 있는 줄을 고친 뒤 다시 올려 주세요.")
            now = datetime.now().isoformat(timespec="seconds")
            con.executemany(
                "INSERT INTO bid_history (bid_date, shipyard, owner, num_ships, quote_usd_mn, market_usd_mn, won,"
                " loss_reason, margin, note, file_name, created_at) VALUES (?,?,?,?,?,?,?,?,?,?,?,?)",
                [(r["date"], r["shipyard"], r["owner"], r["ships"], r["quote"], r["market"], int(r["won"]),
                  r["loss_reason"], r["margin"], r["note"], name, now) for r in new])
            out["saved"] = len(new)
    return out
